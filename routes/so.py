"""
routes/so.py — Special Order document generation routes.
Staff and admin only. Loads so_templates/S_O_TEMP.docx, fills the 5 placeholders,
saves to so_output/, and returns a download link.
"""
import os
import re
from datetime import datetime

from flask import (Blueprint, flash, jsonify, redirect,
                   render_template, request, send_file, session, url_for)

from utils import get_client_ip, login_required

so_bp = Blueprint("so", __name__)

TEMPLATE_PATH = os.path.join("so_templates", "S_O_TEMP.docx")
OUTPUT_DIR    = "so_output"

# ── SO type registry ────────────────────────────────────────────────────────────
# Each entry: label, subject line, list of dynamic field keys, body template.
# Body template uses str.format() — {employee_full_name} and {employee_position}
# are always available; all keys in "fields" must be supplied by the request.

SO_TYPES = {
    "assignment_order": {
        "label":   "Assignment Order",
        "subject": "ASSIGNMENT ORDER",
        "fields":  ["assigned_school", "assigned_district", "assigned_municipality", "effective_date"],
        "body": (
            "In the exigency of the service, you are hereby assigned at "
            "{assigned_school}, {assigned_district}, {assigned_municipality}, Leyte "
            "effective {effective_date} or upon receipt thereof.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "designation_new": {
        "label":   "Designation — New",
        "subject": "DESIGNATION ORDER",
        "fields":  ["designation_role", "school_name", "municipality",
                    "effective_condition", "validity_condition"],
        "body": (
            "In the exigency and best interest of the service, you are hereby designated as "
            "{designation_role} of {school_name}, {municipality}, Leyte effective "
            "{effective_condition}.\n\n"
            "This order is valid {validity_condition}, subject to the request for renewal and "
            "subsequent approval of the Superintendent.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "designation_renewal": {
        "label":   "Designation — Renewal",
        "subject": "RENEWAL OF DESIGNATION",
        "fields":  ["honorific", "designation_role", "school_name", "district", "municipality",
                    "effective_date", "accountability_amount_words",
                    "accountability_amount_figures", "school_year"],
        "body": (
            "With reference to the Approved Plotting for Designation and in the exigency and "
            "best interest of the service, this office renews the designation of "
            "{honorific} {employee_full_name} as {designation_role} of {school_name}, "
            "{district}, {municipality}, Leyte effective {effective_date} or upon receipt "
            "thereof.\n\n"
            "The amount accountability entrusted to the SDO amounting to "
            "{accountability_amount_words} ({accountability_amount_figures}).\n\n"
            "This order is valid for School Year {school_year} only, subject to the request "
            "for renewal and subsequent approval of the Superintendent.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "detailment": {
        "label":   "Detailment",
        "subject": "DETAIL ORDER",
        "fields":  ["honorific", "previous_so_number", "previous_so_year",
                    "from_school", "from_municipality",
                    "to_school", "to_municipality",
                    "effective_date", "school_year"],
        "body": (
            "With reference to the previous Approved Special Order No. {previous_so_number} "
            "s. {previous_so_year} and in the exigency of the service, this office renews the "
            "detail order of {honorific} {employee_full_name} from {from_school}, "
            "{from_municipality}, Leyte to {to_school}, {to_municipality}, Leyte "
            "effective {effective_date}.\n\n"
            "Further, this Order is valid for School Year {school_year}, subject to request "
            "for renewal and subsequent approval of the Superintendent.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "realignment": {
        "label":   "Realignment",
        "subject": "REALIGNMENT",
        "fields":  ["from_school", "from_municipality",
                    "to_school", "to_municipality",
                    "effective_date", "monthly_salary", "nosca_number", "item_note"],
        "body": (
            "Approval is hereby given to the transfer of {employee_position} with "
            "his/her monthly salary, status and effective date:\n\n"
            "Name: {employee_full_name}\n"
            "From: {from_school}, {from_municipality}, LEYTE\n"
            "To: {to_school}, {to_municipality}, LEYTE\n"
            "Effective Date: {effective_date}\n"
            "Salary: P {monthly_salary}\n"
            "Mode of Transfer: REALIGNMENT (NOSCA NO. {nosca_number})\n"
            "Note: {item_note}"
        ),
    },
    "reassignment": {
        "label":   "Reassignment",
        "subject": "REASSIGNMENT",
        "fields":  ["from_school", "from_district", "from_municipality",
                    "to_school", "to_district", "to_municipality",
                    "new_role", "effective_date"],
        "body": (
            "In the exigency and best interest of the service, you are hereby reassigned from "
            "{from_school}, {from_district}, {from_municipality}, Leyte to "
            "{to_school}, {to_district}, {to_municipality}, Leyte as {new_role} "
            "effective {effective_date} or upon receipt thereof.\n\n"
            "Further, this order is subject to the necessary clearance from any money, property "
            "and work-related accountabilities from your previous station before assuming in "
            "your new station.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "resignation": {
        "label":   "Resignation / Termination",
        "subject": "TERMINATION OF SERVICES",
        "fields":  ["school_name", "municipality",
                    "termination_effective_date", "termination_reason"],
        "body": (
            "The services of {employee_full_name}, {employee_position} of "
            "{school_name}, {municipality}, Leyte is hereby TERMINATED effective "
            "{termination_effective_date}, due to {termination_reason}.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "revocation": {
        "label":   "Revocation of Special Order",
        "subject": "REVOCATION OF SPECIAL ORDER",
        "fields":  ["so_number_revoked", "so_year_revoked", "designation_revoked",
                    "school_name", "district", "municipality"],
        "body": (
            "In the exigency of the service, Special Order No. {so_number_revoked} "
            "s. {so_year_revoked} Designated as {designation_revoked} of "
            "{school_name}, {district} {municipality}, Leyte is hereby REVOKED.\n\n"
            "This Order shall take effect immediately.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "shared_services": {
        "label":   "Shared Services",
        "subject": "SHARED SERVICES",
        "fields":  ["assigned_office", "assigned_location",
                    "period_from", "period_to", "days_per_week",
                    "purpose_description", "calendar_year"],
        "body": (
            "In the exigency of the service, you are hereby assigned at the "
            "{assigned_office}, {assigned_location} from {period_from} to {period_to}, "
            "for {days_per_week} days per week.\n\n"
            "Further, you shall assist in the {purpose_description} for Calendar Year "
            "{calendar_year}.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "transfer_of_agency": {
        "label":   "Transfer of Agency",
        "subject": "TRANSFER OF AGENCY",
        "fields":  ["to_agency_name", "to_agency_location", "effective_date"],
        "body": (
            "Approval is hereby given to the transfer of {employee_full_name} "
            "{employee_position} of DepEd Leyte Division to the "
            "{to_agency_name}, {to_agency_location} effective {effective_date}.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "transfer_of_station": {
        "label":   "Transfer of Station",
        "subject": "TRANSFER OF STATION ASSIGNMENT",
        "fields":  ["agreement_basis", "from_school", "from_municipality",
                    "to_school", "to_municipality", "effective_date"],
        "body": (
            "With reference to the {agreement_basis} and in the exigency of the service, "
            "you are hereby reassigned from {from_school}, {from_municipality}, Leyte to "
            "{to_school}, {to_municipality}, Leyte effective on {effective_date}.\n\n"
            "Further, this Order is subject to the necessary clearance from any money, "
            "property, and work-related accountabilities from your previous station before "
            "assuming office in your new station.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
    "transfer_of_station_with_designation": {
        "label":   "Transfer of Station with Designation",
        "subject": "TRANSFER OF STATION ASSIGNMENT WITH DESIGNATION",
        "fields":  ["from_school", "from_district", "from_municipality",
                    "to_district", "to_municipality", "new_role", "effective_date"],
        "body": (
            "In the exigency of the service, you are hereby reassigned from "
            "{from_school}, {from_district}, {from_municipality}, Leyte to "
            "{to_district}, {to_municipality}, Leyte and shall serve as {new_role} of "
            "{to_district}, {to_municipality}, Leyte effective {effective_date} or upon "
            "receipt thereof.\n\n"
            "Further, this Order is subject to the necessary clearance from any money, "
            "property, and work-related accountabilities from your previous station before "
            "assuming office in your new station.\n\n"
            "For your information, proper guidance, and strict compliance."
        ),
    },
}

FIELD_LABELS = {
    "assigned_school":               "Assigned School",
    "assigned_district":             "District",
    "assigned_municipality":         "Municipality",
    "effective_date":                "Effective Date",
    "designation_role":              "Designation Role",
    "school_name":                   "School Name",
    "municipality":                  "Municipality",
    "effective_condition":           "Effective Condition",
    "validity_condition":            "Validity Condition",
    "honorific":                     "Honorific (Ms./Mr./Dr.)",
    "district":                      "District",
    "accountability_amount_words":   "Accountability Amount (in words)",
    "accountability_amount_figures": "Accountability Amount (in figures)",
    "school_year":                   "School Year (e.g. 2026-2027)",
    "previous_so_number":            "Previous SO Number",
    "previous_so_year":              "Previous SO Year",
    "from_school":                   "From School",
    "from_municipality":             "From Municipality",
    "to_school":                     "To School",
    "to_municipality":               "To Municipality",
    "from_district":                 "From District",
    "to_district":                   "To District",
    "new_role":                      "New Role / Designation",
    "monthly_salary":                "Monthly Salary (PHP)",
    "nosca_number":                  "NOSCA Number",
    "item_note":                     "Item Note",
    "termination_effective_date":    "Termination Effective Date",
    "termination_reason":            "Reason for Termination",
    "so_number_revoked":             "SO Number to Revoke",
    "so_year_revoked":               "SO Year",
    "designation_revoked":           "Designation Being Revoked",
    "assigned_office":               "Assigned Office",
    "assigned_location":             "Location",
    "period_from":                   "Period From",
    "period_to":                     "Period To",
    "days_per_week":                 "Days Per Week",
    "purpose_description":           "Purpose / Task Description",
    "calendar_year":                 "Calendar Year",
    "to_agency_name":                "Destination Agency Name",
    "to_agency_location":            "Agency Location",
    "agreement_basis":               "Agreement Basis (e.g. Approved Swapping Agreement)",
}

# Fields that render as date pickers
_DATE_FIELDS = {"effective_date", "termination_effective_date", "period_from", "period_to"}


# ── Date formatter ──────────────────────────────────────────────────────────────

_MONTHS = ["January", "February", "March", "April", "May", "June",
           "July", "August", "September", "October", "November", "December"]


def _fmt_date(s: str) -> str:
    """Convert ISO date '2026-05-11' → 'May 11, 2026'. Passes other strings through."""
    try:
        d = datetime.strptime(s.strip(), "%Y-%m-%d")
        return f"{_MONTHS[d.month - 1]} {d.day}, {d.year}"
    except (ValueError, AttributeError):
        return s


# ── python-docx placeholder helpers ────────────────────────────────────────────

def _replace_in_para_list(paragraphs, placeholder, value):
    """Replace placeholder in a list of paragraphs, handling split runs."""
    for para in paragraphs:
        if placeholder not in para.text:
            continue
        new_text = para.text.replace(placeholder, value)
        runs = para.runs
        if runs:
            runs[0].text = new_text
            for run in runs[1:]:
                run.text = ""
        else:
            para.add_run(new_text)


def _replace_simple(doc, placeholder, value):
    """Replace a single-value placeholder everywhere in the document."""
    _replace_in_para_list(doc.paragraphs, placeholder, value)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_in_para_list(cell.paragraphs, placeholder, value)


def _replace_body(doc, placeholder, body_text):
    """Replace a body placeholder, expanding double-newline blocks into paragraphs."""
    from docx.oxml import OxmlElement
    blocks = [b.strip() for b in body_text.split("\n\n") if b.strip()]
    for para in doc.paragraphs:
        if placeholder not in para.text:
            continue
        parent = para._p.getparent()
        idx    = list(parent).index(para._p)
        for offset, text in enumerate(blocks):
            new_p = OxmlElement("w:p")
            new_r = OxmlElement("w:r")
            new_t = OxmlElement("w:t")
            new_t.text = text
            new_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            new_r.append(new_t)
            new_p.append(new_r)
            parent.insert(idx + offset, new_p)
        parent.remove(para._p)
        break


# ── Auth helper ─────────────────────────────────────────────────────────────────

def _require_staff():
    """Redirect non-staff. Returns a response object or None."""
    if not session.get("logged_in"):
        return redirect(url_for("auth.login"))
    if session.get("role") not in ("staff", "admin"):
        flash("Staff access required.", "error")
        return redirect(url_for("dashboard.index"))
    return None


# ── Routes ──────────────────────────────────────────────────────────────────────

@so_bp.route("/so-request")
@login_required
def so_request_page():
    guard = _require_staff()
    if guard:
        return guard
    return render_template("so_request.html", so_types=SO_TYPES)


@so_bp.route("/api/so/fields/<so_type>")
@login_required
def so_fields(so_type):
    """Return the dynamic field config for a given SO type."""
    if session.get("role") not in ("staff", "admin"):
        return jsonify({"error": "Unauthorized"}), 403
    config = SO_TYPES.get(so_type)
    if not config:
        return jsonify({"error": "Unknown SO type"}), 404
    fields = [
        {
            "key":   f,
            "label": FIELD_LABELS.get(f, f.replace("_", " ").title()),
            "type":  "date" if f in _DATE_FIELDS else "text",
        }
        for f in config["fields"]
    ]
    return jsonify({"fields": fields, "subject": config["subject"]})


@so_bp.route("/api/so/generate", methods=["POST"])
def so_generate():
    """Generate the SO .docx and return a download URL."""
    # Manual auth check — returns JSON so the frontend can display the error
    if not session.get("logged_in") or session.get("role") not in ("staff", "admin"):
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    data               = request.get_json(force=True, silent=True) or {}
    so_type            = data.get("so_type", "").strip()
    employee_full_name = data.get("employee_full_name", "").strip()
    employee_position  = data.get("employee_position", "").strip()
    date_issued        = _fmt_date(data.get("date_issued", "").strip())
    fields             = {k: v.strip() for k, v in data.get("fields", {}).items()}

    # ── Validate ────────────────────────────────────────────────────────────────
    if not so_type or so_type not in SO_TYPES:
        return jsonify({"success": False, "message": "Invalid SO type"}), 400
    if not employee_full_name:
        return jsonify({"success": False, "message": "Employee name is required"}), 400
    if not employee_position:
        return jsonify({"success": False, "message": "Employee position is required"}), 400
    if not date_issued:
        return jsonify({"success": False, "message": "Date issued is required"}), 400

    config = SO_TYPES[so_type]

    # Format any date-type dynamic fields
    for key in _DATE_FIELDS:
        if key in fields:
            fields[key] = _fmt_date(fields[key])

    # ── Build body ──────────────────────────────────────────────────────────────
    body_vars = {
        "employee_full_name": employee_full_name,
        "employee_position":  employee_position,
        **fields,
    }
    try:
        body = config["body"].format(**body_vars)
    except KeyError as missing:
        return jsonify({"success": False, "message": f"Missing required field: {missing}"}), 400

    # ── Load template ───────────────────────────────────────────────────────────
    if not os.path.exists(TEMPLATE_PATH):
        return jsonify({"success": False,
                        "message": "SO template not found. Place S_O_TEMP.docx in so_templates/."}), 500

    try:
        from docx import Document
        doc = Document(TEMPLATE_PATH)

        # Replace body FIRST (it expands into multiple paragraphs)
        _replace_body(doc, "{{body}}", body)

        # Replace the four single-value placeholders
        _replace_simple(doc, "{{employee_full_name}}", employee_full_name)
        _replace_simple(doc, "{{employee_position}}",  employee_position)
        _replace_simple(doc, "{{subject}}",            config["subject"])
        _replace_simple(doc, "{{date_issued}}",        date_issued)

        # ── Save ─────────────────────────────────────────────────────────────────
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"SO_{so_type.upper()}_{ts}.docx"
        doc.save(os.path.join(OUTPUT_DIR, filename))

        try:
            from services.misc import audit_log
            audit_log(
                "so_generated",
                f"so_type={so_type} employee={employee_full_name} file={filename}",
                username=session.get("username", "?"),
                ip=get_client_ip(),
            )
        except Exception:
            pass

        return jsonify({
            "success":      True,
            "filename":     filename,
            "download_url": url_for("so.so_download", filename=filename),
        })

    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@so_bp.route("/api/so/download/<filename>")
@login_required
def so_download(filename):
    """Serve the generated .docx file for download."""
    if session.get("role") not in ("staff", "admin"):
        return jsonify({"error": "Unauthorized"}), 403

    # Only allow filenames that match the generator's pattern
    if not re.match(r"^SO_[A-Z0-9_]+\.docx$", filename):
        return jsonify({"error": "Invalid filename"}), 400

    filepath = os.path.join(OUTPUT_DIR, filename)
    if not os.path.exists(filepath):
        return jsonify({"error": "File not found"}), 404

    return send_file(
        filepath,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )
